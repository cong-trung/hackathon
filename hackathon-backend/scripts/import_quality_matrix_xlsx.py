import argparse
import csv
import io
import re
import shutil
import sys
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.io.concurrency import atomic_write_bytes, locked_files

IMPORTER = "quality_matrix_xlsx_import"
SEED_MARKERS = {"product_catalog_seed", "seed_script"}
CELL_HEADERS = [
    "submission_id", "question_id", "module_key", "role", "answer",
    "updated_at", "updated_by",
]
SUBMISSION_HEADERS = [
    "submission_id", "created_at", "updated_at", "submitted_by", "role",
    "selected_modules", "product_name", "prodgroup3", "operation",
    "answer_count", "status",
]
FIELD_MAP = {
    "product_name": "productsname",
    "prodgroup3": "prodgroup3",
    "tp_name": "testprogramnamefromtriggeroperation",
    "operation_input": "maintestprocessflowoperation",
    "backend_test_process_flow": "backendtestprocessflow",
    "nsb": "nsbnonsynergybuild",
    "mcp": "mcpmultichippackageproductshavemultipledies",
    "package": "packagebgalgapgafoveros",
    "bare_die_lid": "baredieorlidproduct",
    "edm_check_enable": "edmcheckenableindcfailflow",
    "implementation_stage": "implementationstage",
    "segment": "productsegment",
    "xvi_tool_type": "xvitooltype",
    "fab_technology": "fabtechnology",
    "hdmx_ap_pan": "hdmxappan",
    "lts_ball": "ltsballproducsyn",
    "ppv_sw_platform": "ppvswplatform",
    "thermal_technology": "ppvthermaltechnology",
}
QUESTION_IDS = {
    "product_name", "prodgroup3", "tp_name", "operation", "process_step", "foveros",
    "backend_test_process_flow", "nsb", "mcp", "package", "bare_die_lid",
    "edm_check_enable", "implementation_stage", "segment", "xvi_tool_type",
    "fab_technology", "hdmx_ap_pan", "lts_ball", "ppv_sw_platform",
    "thermal_technology",
}
Y_FIELDS = {"nsb", "mcp", "edm_check_enable", "lts_ball"}
BARE_DIE_ALIASES = {
    "bae die": "Bare die", "baredie": "Bare die", "bare die": "Bare die",
    "lid product": "Lidded", "lid": "Lidded", "lidded product": "Lidded",
}
PPV_FIELDS = {"ppv_sw_platform", "thermal_technology"}
EXPECTED_FILE_HEADERS = {"submissions": SUBMISSION_HEADERS, "cells": CELL_HEADERS}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def read_csv(path):
    raw = path.read_bytes()
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        return raw, reader.fieldnames or [], list(reader)


def encode_csv(headers, rows, original_bytes):
    output = io.StringIO(newline="")
    line_ending = "\r\n" if b"\r\n" in original_bytes else "\n"
    writer = csv.DictWriter(
        output, fieldnames=headers, extrasaction="ignore", lineterminator=line_ending
    )
    writer.writeheader()
    writer.writerows(
        {header: "" if row.get(header) is None else str(row.get(header, "")) for header in headers}
        for row in rows
    )
    encoded = output.getvalue().encode("utf-8")
    return (b"\xef\xbb\xbf" if original_bytes.startswith(b"\xef\xbb\xbf") else b"") + encoded


def ensure_local_paths(data_dir, xlsx):
    data_dir = data_dir.expanduser().resolve()
    xlsx = xlsx.expanduser().resolve()
    if not data_dir.is_dir():
        raise ValueError(f"--data-dir is not a directory: {data_dir}")
    expected_files = (
        Path("QualityMatrix/PYPCS_Submissions.csv"),
        Path("QualityMatrix/PYPCS_SubmissionCells.csv"),
        Path("QualityMatrix/Settings/PYPCS_Questions.csv"),
        Path("QualityMatrix/Settings/Operation_Process_Step.csv"),
        Path("QualityMatrix/Settings/PYPCS_QuestionModuleMap.csv"),
        Path("QualityMatrix/Settings/PYPCS_Modules.csv"),
    )
    missing_files = [str(path) for path in expected_files if not (data_dir / path).is_file()]
    if missing_files:
        raise ValueError(f"--data-dir is missing expected layout files: {missing_files}")
    if not xlsx.is_file() or xlsx.suffix.casefold() != ".xlsx":
        raise ValueError(f"--xlsx must be an existing .xlsx file: {xlsx}")
    return data_dir, xlsx


def workbook_columns(header_row):
    normalized = {normalized_header(value): index for index, value in enumerate(header_row) if value}
    columns = {}
    for field, expected in FIELD_MAP.items():
        matches = [index for header, index in normalized.items() if header.startswith(expected) or expected in header]
        if len(matches) != 1:
            raise ValueError(f"Expected exactly one workbook header for {field}; found {len(matches)}")
        columns[field] = matches[0]
    dropped = {
        "npi_hvm": ("npihvm", "npi or hvm"),
        "new_finish_flow": ("anynewinfinishtestflow",),
        "prq_change": ("anychangestowardsprqperiod",),
        "request": ("request",),
    }
    headers = list(normalized)
    for field, candidates in dropped.items():
        if not any(candidate.replace(" ", "") in header for candidate in candidates for header in headers):
            raise ValueError(f"Expected workbook header for dropped field {field} is missing")
    return columns


def module_for_label(label):
    text = normalized_header(label)
    if "ppv" in text or "sthi" in text or text in {"sps", "spspm"}:
        return "STHI"
    if "lcbi" in text:
        return "LCBI"
    if "hdbi" in text or text.startswith("hbi"):
        return "HDBI"
    if "olb" in text:
        return "OLB"
    if "v8" in text:
        return "V8"
    if "ahmt" in text:
        return "HDMX_AHMT"
    if text == "hdmx" or text.startswith("hdmxonly"):
        return "HDMX"
    if any(label in text for label in ("phmhot", "phmcold", "phmcolr", "phmhotr")):
        return "HDMx_EQA"
    if text == "csm" or text.startswith(("classhot", "classcold")):
        return "HDMx_Class"
    if "class" in text:
        return "HDMx_Class"
    if "cfuse" in text or "fuse" in text:
        return "HDMx_Cfuse"
    if "eqa" in text or text.startswith("eq"):
        return "HDMx_EQA"
    if "rebi" in text:
        return "REBI"
    if "klot" in text:
        return "KLOT"
    return ""


def operation_index(operation_rows):
    owners = defaultdict(set)
    codes = set()
    for row in operation_rows:
        code = clean(row.get("OPERATION"))
        module = clean(row.get("MODULE_KEY"))
        if code and module:
            codes.add(code)
            owners[code.casefold()].add(module)
    pattern = re.compile(
        r"(?<![A-Za-z0-9])(" + "|".join(re.escape(code) for code in sorted(codes, key=len, reverse=True)) + r")(?![A-Za-z0-9])",
        re.IGNORECASE,
    )
    return owners, pattern


def _label_module(label, position, text):
    module = module_for_label(label)
    if module == "REBI" or module == "KLOT":
        candidates = list(re.finditer(r"\b(LCBI|HDBI)\b", text, re.IGNORECASE))
        preceding = [match for match in candidates if match.start() <= position]
        nearest = max(preceding, key=lambda match: match.start()) if preceding else min(
            candidates, key=lambda match: match.start(), default=None
        )
        return nearest.group(1).upper() if nearest else ""
    if module == "HDMX":
        return ""
    if "fuse" in normalized_header(label) and "cfuse" not in normalized_header(label):
        line_start = text.rfind("\n", 0, position) + 1
        line_end = text.find("\n", position)
        if line_end < 0:
            line_end = len(text)
        nearby = list(re.finditer(r"\b(OLB|V8)\b", text[line_start:line_end], re.IGNORECASE))
        if nearby:
            local_position = position - line_start
            preceding = [match for match in nearby if match.start() <= local_position]
            nearest = max(preceding, key=lambda match: match.start()) if preceding else min(
                nearby, key=lambda match: match.start()
            )
            return nearest.group(1).upper()
    return module


def _short_snippet(text, start, end, limit=180):
    snippet = re.sub(r"\s+", " ", text[max(0, start - 55):min(len(text), end + 85)]).strip()
    return snippet if len(snippet) <= limit else snippet[:limit - 3] + "..."


def parse_operations(text, owners, pattern, module_keys_by_upper, catalog_pairs, row_number):
    answers = defaultdict(list)
    steps = defaultdict(list)
    unmatched = Counter()
    custom_values = Counter()
    flags = set()
    additions = defaultdict(set)
    expert_evidence = []
    matched_tokens = 0
    candidate_tokens = 0
    module_context = ""
    step_context = ""
    if text.strip().casefold() == "dln":
        unmatched["dln"] += 1
        flags.add("dropped dln")
        return answers, steps, unmatched, custom_values, additions, flags, expert_evidence, matched_tokens, candidate_tokens

    numeric_pattern = re.compile(r"(?<![A-Za-z0-9])\d{3,6}(?![A-Za-z0-9])")
    boundaries = list(re.finditer(r"\n|;|->", text))
    starts = [0, *(match.end() for match in boundaries)]
    ends = [*(match.start() for match in boundaries), len(text)]
    for start, end in zip(starts, ends):
        segment = text[start:end]
        numeric_matches = [
            match for match in numeric_pattern.finditer(segment)
            if len(match.group(0)) >= 4 or match.group(0).casefold() in owners
        ]
        if not numeric_matches:
            continue
        first_position = numeric_matches[0].start()
        colon = segment.find(":")
        explicit_label = colon >= 0 and colon < first_position
        label = segment[:colon].strip().strip('"') if explicit_label else segment[:first_position].strip(" \t-_,/()")
        label_module = module_for_label(label) if label else ""
        simple_unplaced_prefix = bool(
            label
            and re.fullmatch(r"[A-Za-z][A-Za-z0-9 ()/-]{0,40}", label)
            and not label_module
        )
        has_label = bool(label and (explicit_label or label_module or simple_unplaced_prefix))
        segment_module = ""
        if has_label:
            segment_module = _label_module(label, start, text)
            step_context = label
            if segment_module:
                module_context = segment_module
            else:
                module_context = ""
                normalized_label = normalized_header(label)
                if label_module in {"REBI", "KLOT"}:
                    flags.add(f"unresolved {label.strip()}")
                elif label_module == "HDMX":
                    flags.add("bare HDMx")
                else:
                    flags.add(f"unplaced label {label.strip()}")

        unknown_label = has_label and not segment_module
        if unknown_label:
            codes = [match.group(0) for match in numeric_matches]
            expert_evidence.append({
                "label": label.strip(),
                "row": row_number,
                "snippet": _short_snippet(text, start, end),
                "codes": codes,
            })

        candidate_tokens += len(numeric_matches)
        for match in numeric_matches:
            code = match.group(0)
            possible_owners = owners.get(code.casefold(), set())
            if unknown_label:
                unmatched[code] += 1
                continue
            assigned_module = segment_module or module_context
            if assigned_module:
                assigned_modules = [assigned_module]
            elif len(possible_owners) == 1:
                assigned_modules = [next(iter(possible_owners))]
            elif len(possible_owners) > 1:
                unmatched[code] += 1
                flags.add(f"multi-owner unlabeled {code}")
                continue
            else:
                unmatched[code] += 1
                flags.add(f"unassigned code {code}")
                continue

            matched_tokens += 1
            for module in assigned_modules:
                canonical_module = module_keys_by_upper.get(module.upper())
                if not canonical_module:
                    unmatched[code] += 1
                    flags.add(f"unknown module {module}")
                    continue
                if code not in answers[canonical_module]:
                    answers[canonical_module].append(code)
                if (canonical_module.upper(), code.casefold()) not in catalog_pairs:
                    additions[(canonical_module, code)].add(row_number)
                if (canonical_module.upper(), code.casefold()) not in catalog_pairs:
                    custom_values[("operation", code)] += 1
                if step_context and step_context not in steps[canonical_module]:
                    steps[canonical_module].append(step_context)
    if text.strip() and text.strip().casefold() not in {"na", "n/a", "none"} and not candidate_tokens:
        unmatched[text.strip()] += 1
        flags.add("unassigned operation text")
    return answers, steps, unmatched, custom_values, additions, flags, expert_evidence, matched_tokens, candidate_tokens


def normalize_value(question_id, raw_value, question, changes, other_values, na_values):
    value = clean(raw_value)
    if not value:
        return ""
    original = value
    if question_id in Y_FIELDS:
        folded = value.casefold()
        if folded in {"yes", "y"}:
            value = "Y"
        elif folded in {"no", "n", "not nsb", "non-nsb", "non nsb"}:
            value = "N"
        elif folded in {"na", "n/a"}:
            na_values[(question_id, original)] += 1
            value = ""
    elif question_id == "bare_die_lid":
        value = BARE_DIE_ALIASES.get(value.casefold(), value)
    elif question_id == "package":
        has_bga = "bga" in value.casefold()
        has_lga = "lga" in value.casefold()
        if has_bga and has_lga:
            pass
        elif has_bga:
            value = "BGA"
        elif has_lga:
            value = "LGA"
    if value and question.get("field_type", "").casefold() == "select":
        options = [item.strip() for item in question.get("options", "").split("|") if item.strip()]
        option = next((item for item in options if item.casefold() == value.casefold()), None)
        if option:
            value = option
        elif question_id != "prodgroup3":
            other_values[(question_id, value)] += 1
    if original != value:
        changes[(question_id, original, value)] += 1
    return value


def create_plan(data_dir, xlsx_path):
    quality_dir = data_dir / "QualityMatrix"
    settings = quality_dir / "Settings"
    questions_path = settings / "PYPCS_Questions.csv"
    submissions_path = quality_dir / "PYPCS_Submissions.csv"
    cells_path = quality_dir / "PYPCS_SubmissionCells.csv"
    question_bytes, question_headers, questions = read_csv(questions_path)
    submission_bytes, submission_headers, submissions = read_csv(submissions_path)
    cell_bytes, cell_headers, cells = read_csv(cells_path)
    for label, headers in (("submissions", submission_headers), ("cells", cell_headers)):
        missing = [header for header in EXPECTED_FILE_HEADERS[label] if header not in headers]
        if missing:
            raise ValueError(f"{label} CSV is missing expected columns: {missing}")
    questions_by_id = {row.get("question_id", ""): row for row in questions}
    missing_questions = QUESTION_IDS.difference(questions_by_id)
    if missing_questions:
        raise ValueError(f"PYPCS_Questions.csv is missing question IDs: {sorted(missing_questions)}")
    map_path = settings / "PYPCS_QuestionModuleMap.csv"
    with map_path.open(encoding="utf-8-sig", newline="") as source:
        map_rows = list(csv.DictReader(source))
    applicable = {
        (clean(row.get("question_id")), clean(row.get("module_key")).upper())
        for row in map_rows
        if clean(row.get("visible_default")).upper() in {"Y", "YES", "TRUE", "1"}
    }
    operation_path = settings / "Operation_Process_Step.csv"
    operation_bytes, operation_headers, operation_rows = read_csv(operation_path)
    missing_operation_headers = {
        "MODULE_KEY", "OPERATION", "OPER_SHORT_DESC", "OPER_LONG_DESC"
    }.difference(operation_headers)
    if missing_operation_headers:
        raise ValueError(
            f"Operation_Process_Step.csv is missing columns: {sorted(missing_operation_headers)}"
        )
    with (settings / "PYPCS_Modules.csv").open(encoding="utf-8-sig", newline="") as source:
        module_rows = list(csv.DictReader(source))
    module_order = list(dict.fromkeys(clean(row.get("module_key")) for row in module_rows if clean(row.get("module_key"))))
    question_order = [row["question_id"] for row in questions if row.get("question_id") in QUESTION_IDS]

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_name = "New NPI request (TI's update)"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = workbook_columns(header_row)
    module_keys_by_upper = {module.upper(): module for module in module_order}
    catalog_pairs = {
        (clean(row.get("MODULE_KEY")).upper(), clean(row.get("OPERATION")).casefold())
        for row in operation_rows
        if clean(row.get("MODULE_KEY")) and clean(row.get("OPERATION"))
    }
    owners, operation_pattern = operation_index(operation_rows)
    operation_additions = defaultdict(set)
    source_rows = []
    skipped = []
    normalization_changes = Counter()
    other_values = Counter()
    na_values = Counter()
    unmatched_operations = Counter()
    fallback_rows = []
    expert_evidence = []
    previous_importer = [row for row in submissions if clean(row.get("submitted_by")).casefold() == IMPORTER]
    previous_by_product = defaultdict(list)
    for row in previous_importer:
        key = (clean(row.get("prodgroup3")).casefold(), clean(row.get("product_name")).casefold())
        previous_by_product[key].append(clean(row.get("selected_modules")).split("|"))
    row_module_changes = []
    module_distribution = Counter()
    operation_match_count = 0
    operation_candidate_count = 0
    package_source_distribution = Counter()
    blank_package_rows = []
    foveros_values = Counter()
    pg_options = [item.strip() for item in questions_by_id["prodgroup3"].get("options", "").split("|") if item.strip()]
    pg_seen = {item.casefold() for item in pg_options}
    new_pg_options = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = {field: clean(row[index] if index < len(row) else "") for field, index in columns.items()}
        if row_number == 44:
            skipped.append((row_number, "excluded aggregate row: All products at STHI"))
            continue
        if normalized_header(raw["product_name"]) == "productsname" or normalized_header(raw["prodgroup3"]) == "prodgroup3":
            skipped.append((row_number, "header echo"))
            continue
        reasons = []
        if not raw["product_name"]:
            reasons.append("blank product name")
        if not raw["prodgroup3"]:
            reasons.append("blank prodgroup3")
        if reasons:
            skipped.append((row_number, " and ".join(reasons)))
            continue
        package_source_distribution[raw["package"] or "<blank>"] += 1
        if not raw["package"]:
            blank_package_rows.append(row_number)
        pg = raw["prodgroup3"]
        if pg.casefold() not in pg_seen:
            pg_seen.add(pg.casefold())
            new_pg_options.append(pg)
        normalized = {}
        for field in FIELD_MAP:
            if field == "operation_input":
                continue
            normalized[field] = normalize_value(
                field, raw[field], questions_by_id[field], normalization_changes, other_values, na_values
            )
        package_text = raw["package"].casefold()
        if "non-foveros" in package_text or "non foveros" in package_text:
            normalized["foveros"] = "Non-Foveros"
        elif "foveros" in package_text:
            normalized["foveros"] = "Foveros"
        if normalized.get("foveros"):
            foveros_values[normalized["foveros"]] += 1
        operation_values, step_values, unmatched, custom_ops, additions, flags, row_evidence, matched, candidates = parse_operations(
            raw["operation_input"], owners, operation_pattern, module_keys_by_upper,
            catalog_pairs, row_number,
        )
        for key, source_numbers in additions.items():
            operation_additions[key].update(source_numbers)
        expert_evidence.extend(row_evidence)
        operation_match_count += matched
        operation_candidate_count += candidates
        unmatched_operations.update(unmatched)
        other_values.update(custom_ops)
        if any(normalized.get(field, "") and normalized[field].casefold() not in {"na", "n/a"} for field in PPV_FIELDS):
            operation_values.setdefault("STHI", [])
        modules = [module for module in module_order if module in operation_values or module in step_values]
        if not modules:
            modules = ["STHI"]
            fallback_rows.append(row_number)
            flags.add("fallback STHI")
        for module in modules:
            module_distribution[module] += 1
        mapped_values = dict(normalized)
        mapped_values["operation"] = {module: "|".join(operation_values.get(module, [])) for module in modules}
        mapped_values["process_step"] = {module: "|".join(step_values.get(module, [])) for module in modules}
        source_rows.append({
            "source_row": row_number,
            "product_name": normalized["product_name"],
            "prodgroup3": normalized["prodgroup3"],
            "modules": modules,
            "values": mapped_values,
            "flags": sorted(flags),
        })
        prior_key = (normalized["prodgroup3"].strip().casefold(), normalized["product_name"].strip().casefold())
        prior_modules = previous_by_product[prior_key].pop(0) if previous_by_product[prior_key] else None
        if prior_modules is not None and {module.upper() for module in prior_modules if module} != {module.upper() for module in modules}:
            row_module_changes.append({"row": row_number, "before": prior_modules, "after": modules})
    workbook.close()

    new_operation_rows = [
        {
            "MODULE_KEY": module,
            "OPERATION": operation,
            "OPER_SHORT_DESC": "",
            "OPER_LONG_DESC": "Added from Test Quality Product Matrix import",
        }
        for module, operation in sorted(
            operation_additions,
            key=lambda item: (item[0].casefold(), int(item[1]) if item[1].isdigit() else item[1]),
        )
    ]
    updated_operation_rows = [*operation_rows, *new_operation_rows]

    groups = defaultdict(list)
    for item in source_rows:
        key = f"{item['prodgroup3'].strip().casefold()}||{item['product_name'].strip().casefold()}"
        groups[key].append(item)
    duplicate_groups = []
    products = []
    for key, items in groups.items():
        if len(items) == 1:
            products.append(items[0])
            continue
        signatures = {
            (tuple(item["modules"]), tuple(sorted((field, repr(value)) for field, value in item["values"].items())))
            for item in items
        }
        identical = len(signatures) == 1
        duplicate_groups.append({
            "key": key,
            "rows": [item["source_row"] for item in items],
            "handling": "import once (identical mapped values)" if identical else "import every row (conflicting mapped values)",
        })
        products.extend(items[:1] if identical else items)

    current_importer = previous_importer
    marker_submissions = [
        row for row in submissions
        if clean(row.get("submitted_by")).casefold() in SEED_MARKERS
    ]
    cells_by_submission = defaultdict(list)
    for row in cells:
        cells_by_submission[clean(row.get("submission_id"))].append(row)
    seeds = []
    retained_seed_markers = []
    for row in marker_submissions:
        marker = clean(row.get("submitted_by")).casefold()
        submission_id = clean(row.get("submission_id"))
        submission_cells = cells_by_submission[submission_id]
        mismatched_cells = [
            cell for cell in submission_cells
            if clean(cell.get("updated_by")).casefold() != marker
        ]
        if not clean(row.get("updated_at")) and not mismatched_cells:
            seeds.append(row)
        else:
            reasons = []
            if clean(row.get("updated_at")):
                reasons.append("updated_at is not blank")
            if mismatched_cells:
                reasons.append("one or more cells have a different updated_by")
            retained_seed_markers.append((submission_id, "; ".join(reasons)))
    user_submissions = [
        row for row in submissions
        if clean(row.get("submitted_by")).casefold() not in {IMPORTER, *SEED_MARKERS}
    ]
    remove_ids = {clean(row.get("submission_id")) for row in [*current_importer, *seeds]}
    retained_submissions = [row for row in submissions if clean(row.get("submission_id")) not in remove_ids]
    retained_cells = [row for row in cells if clean(row.get("submission_id")) not in remove_ids]
    now = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    new_submissions = []
    new_cells = []
    for item in products:
        submission_id = str(uuid.uuid4())
        module_keys = item["modules"]
        for question_id in question_order:
            for module in module_keys:
                if (question_id, module.upper()) not in applicable:
                    continue
                question = questions_by_id[question_id]
                value = item["values"].get(question_id, "")
                if isinstance(value, dict):
                    value = value.get(module, "")
                if not value:
                    continue
                new_cells.append({
                    "submission_id": submission_id,
                    "question_id": question_id,
                    "module_key": module,
                    "role": clean(question.get("role")),
                    "answer": value,
                    "updated_at": "",
                    "updated_by": IMPORTER,
                })
        op_value = next((item["values"]["operation"].get(module, "") for module in module_keys if item["values"]["operation"].get(module, "")), "")
        new_submissions.append({
            "submission_id": submission_id,
            "created_at": now,
            "updated_at": "",
            "submitted_by": IMPORTER,
            "role": "TI",
            "selected_modules": "|".join(module_keys),
            "product_name": item["product_name"],
            "prodgroup3": item["prodgroup3"],
            "operation": op_value,
            "answer_count": str(sum(row["submission_id"] == submission_id for row in new_cells)),
            "status": "draft",
        })

    def package_distribution(cell_rows, submission_ids=None):
        distribution = defaultdict(Counter)
        for cell in cell_rows:
            if clean(cell.get("question_id")) != "package":
                continue
            if submission_ids is not None and clean(cell.get("submission_id")) not in submission_ids:
                continue
            distribution[clean(cell.get("module_key"))][clean(cell.get("answer"))] += 1
        return distribution

    previous_importer_ids = {clean(row.get("submission_id")) for row in previous_importer}
    new_importer_ids = {clean(row.get("submission_id")) for row in new_submissions}
    package_before = package_distribution(cells)
    package_importer_before = package_distribution(cells, previous_importer_ids)
    package_after = package_distribution(retained_cells + new_cells)
    package_importer_after = package_distribution(new_cells, new_importer_ids)

    updated_questions = [dict(row) for row in questions]
    for row in updated_questions:
        if row.get("question_id") == "prodgroup3":
            options = [item.strip() for item in row.get("options", "").split("|") if item.strip()]
            non_other = []
            seen = set()
            for item in options:
                if item.casefold() != "other" and item.casefold() not in seen:
                    seen.add(item.casefold())
                    non_other.append(item)
            for item in new_pg_options:
                if item.casefold() not in seen:
                    seen.add(item.casefold())
                    non_other.append(item)
            row["options"] = "|".join([*non_other, "Other"])
            break
    return {
        "data_dir": data_dir,
        "xlsx_path": xlsx_path,
        "paths": (submissions_path, cells_path, questions_path, operation_path),
        "originals": (
            (submission_bytes, submission_headers),
            (cell_bytes, cell_headers),
            (question_bytes, question_headers),
            (operation_bytes, operation_headers),
        ),
        "submissions": retained_submissions + new_submissions,
        "cells": retained_cells + new_cells,
        "questions": updated_questions,
        "source_rows": source_rows,
        "skipped": skipped,
        "products": products,
        "duplicate_groups": duplicate_groups,
        "seed_count": len(seeds),
        "retained_seed_markers": retained_seed_markers,
        "importer_count": len(current_importer),
        "user_count": len(user_submissions),
        "module_distribution": module_distribution,
        "normalization_changes": normalization_changes,
        "other_values": other_values,
        "na_values": na_values,
        "unmatched_operations": unmatched_operations,
        "fallback_rows": fallback_rows,
        "expert_evidence": expert_evidence,
        "row_module_changes": row_module_changes,
        "package_source_distribution": package_source_distribution,
        "blank_package_rows": blank_package_rows,
        "package_before": package_before,
        "package_after": package_after,
        "package_importer_before": package_importer_before,
        "package_importer_after": package_importer_after,
        "foveros_values": foveros_values,
        "operation_match_count": operation_match_count,
        "operation_candidate_count": operation_candidate_count,
        "new_pg_options": new_pg_options,
        "operation_additions": operation_additions,
        "new_operation_rows": new_operation_rows,
        "operation_rows": updated_operation_rows,
        "new_submissions": new_submissions,
        "new_cells": new_cells,
    }


def print_preview(plan):
    print(f"Workbook rows: {len(plan['source_rows']) + len(plan['skipped'])}; products: {len(plan['source_rows'])}; skipped: {len(plan['skipped'])}")
    for row_number, reason in plan["skipped"]:
        print(f"  SKIP row {row_number}: {reason}")
    print(f"Products to create: {len(plan['products'])}; seed submissions to remove: {plan['seed_count']}; prior importer rows to replace: {plan['importer_count']}; user submissions kept: {plan['user_count']}")
    print(f"Seed-marker submissions retained by safety check: {len(plan['retained_seed_markers'])}")
    for submission_id, reason in plan["retained_seed_markers"]:
        print(f"  submission_id={submission_id}: {reason}")
    print(f"Duplicate groups: {len(plan['duplicate_groups'])}")
    for group in plan["duplicate_groups"]:
        print(f"  {group['key']} rows {group['rows']}: {group['handling']}")
    print("Modules distribution: " + ", ".join(f"{key}={value}" for key, value in sorted(plan["module_distribution"].items())))
    print("Package source distribution: " + ", ".join(
        f"{value!r}={count}" for value, count in sorted(plan["package_source_distribution"].items())
    ))
    print(f"Genuinely blank package source rows: {plan['blank_package_rows']}")
    for label in ("before", "after", "importer_before", "importer_after"):
        distribution = plan[f"package_{label}"]
        modules = "; ".join(
            f"{module}[{', '.join(f'{value!r}={count}' for value, count in sorted(values.items()))}]"
            for module, values in sorted(distribution.items())
        )
        print(f"Package cells {label}: {modules or 'none'}")
    print(f"Foveros values set: {dict(plan['foveros_values']) or 'none'}")
    print("Per-row module assignments:")
    for item in plan["source_rows"]:
        flags = "; ".join(item["flags"]) or "-"
        product_name = re.sub(r"\s+", " ", item["product_name"]).strip()
        prodgroup3 = re.sub(r"\s+", " ", item["prodgroup3"]).strip()
        print(
            f"  row={item['source_row']} prodgroup3={prodgroup3} "
            f"product={product_name} modules={'|'.join(item['modules'])} flags={flags}"
        )
    flags_by_reason = defaultdict(list)
    for item in plan["source_rows"]:
        for flag in item["flags"]:
            flags_by_reason[flag].append(item["source_row"])
    print("Remaining flags by reason:")
    for flag, rows in sorted(flags_by_reason.items()):
        print(f"  {flag}: rows={rows}")
    print("Module changes from previous import:")
    if plan["row_module_changes"]:
        for change in plan["row_module_changes"]:
            print(f"  row={change['row']} before={'|'.join(change['before'])} after={'|'.join(change['after'])}")
    else:
        print("  none")
    rate = plan["operation_match_count"] / plan["operation_candidate_count"] if plan["operation_candidate_count"] else 1.0
    print(f"Operation codes matched: {plan['operation_match_count']}/{plan['operation_candidate_count']} ({rate:.1%})")
    print(f"Operations to add to Operation_Process_Step.csv: {len(plan['new_operation_rows'])}")
    modules_by_operation = defaultdict(set)
    for module, operation in plan["operation_additions"]:
        modules_by_operation[operation].add(module)
    for row in plan["new_operation_rows"]:
        key = (row["MODULE_KEY"], row["OPERATION"])
        source_rows = sorted(plan["operation_additions"][key])
        print(
            f"  module={row['MODULE_KEY']} operation={row['OPERATION']} "
            f"source rows={source_rows}"
        )
    conflicts = {
        operation: sorted(modules)
        for operation, modules in modules_by_operation.items()
        if len(modules) > 1
    }
    if conflicts:
        print("Operation module conflicts:")
        for operation, modules in sorted(conflicts.items()):
            print(f"  operation={operation} modules={modules}")
    print("Normalization changes:")
    for (field, before, after), count in sorted(plan["normalization_changes"].items()):
        print(f"  {field}: {before!r} -> {after!r} ({count})")
    print("NA values left blank:")
    for (field, value), count in sorted(plan["na_values"].items()):
        print(f"  {field}: {value!r} ({count})")
    print("Values stored using Other custom text:")
    other_by_field = defaultdict(Counter)
    for (field, value), count in plan["other_values"].items():
        other_by_field[field][value] += count
    for field, values in sorted(other_by_field.items()):
        print(f"  {field}: {sum(values.values())} values, {len(values)} unique")
        for value, count in values.most_common(5):
            print(f"    {value!r} ({count})")
        if len(values) > 5:
            print(f"    ... {len(values) - 5} additional unique values")
    print("Unmatched operation tokens (top 20):")
    for token, count in plan["unmatched_operations"].most_common(20):
        print(f"  {token!r}: {count}")
    print(f"Rows falling back to STHI: {plan['fallback_rows']}")
    print("Unplaced-label expert evidence:")
    catalog_by_code = defaultdict(list)
    for operation_row in plan["operation_rows"]:
        catalog_by_code[clean(operation_row.get("OPERATION")).casefold()].append(operation_row)
    suggestions = {
        "qahot": "QA hot",
        "qaextremecold": "QA extreme cold",
        "fcb40c": "FCB (-40C)",
        "sh2bioperation": "SH2BI operation",
        "ast": "AST (confirm)",
        "plts": "PLTS (confirm)",
        "bi": "BI (confirm meaning)",
    }
    for evidence in plan["expert_evidence"]:
        code_details = []
        for code in evidence["codes"]:
            catalog_rows = catalog_by_code.get(code.casefold(), [])
            if not catalog_rows:
                code_details.append(f"{code}: not in catalog")
                continue
            for catalog_row in catalog_rows:
                code_details.append(
                    f"{code}: {catalog_row['MODULE_KEY']} "
                    f"(short={catalog_row['OPER_SHORT_DESC']!r}, long={catalog_row['OPER_LONG_DESC']!r})"
                )
        label_key = normalized_header(evidence["label"])
        suggestion = suggestions.get(label_key, f"{evidence['label']} (confirm)")
        print(
            f"  {evidence['label']} | suggestion={suggestion} | rows=[{evidence['row']}] "
            f"| snippet={evidence['snippet']!r} | codes={'; '.join(code_details) or 'none'}"
        )
    print(f"New prodgroup3 options added: {len(plan['new_pg_options'])}: {plan['new_pg_options']}")
    print(f"New cells: {len(plan['new_cells'])}")


def apply_plan(plan):
    data_dir = plan["data_dir"]
    paths = plan["paths"]
    with locked_files(*paths):
        plan = create_plan(data_dir, plan["xlsx_path"])
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        backup_dir = data_dir / "backups" / f"quality_matrix_import_{timestamp}"
        if backup_dir.exists():
            raise FileExistsError(f"Backup directory already exists: {backup_dir}")
        paths = plan["paths"]
        raw_files = [path.read_bytes() for path in paths]
        backup_dir.mkdir(parents=True)
        for path, raw in zip(paths, raw_files):
            backup_path = backup_dir / path.name
            shutil.copy2(path, backup_path)
            if backup_path.read_bytes() != raw:
                raise RuntimeError(f"Backup verification failed for {path.name}")
        (
            (submission_original, submission_headers),
            (cell_original, cell_headers),
            (question_original, question_headers),
            (operation_original, operation_headers),
        ) = plan["originals"]
        payloads = (
            encode_csv(submission_headers, plan["submissions"], submission_original),
            encode_csv(cell_headers, plan["cells"], cell_original),
            encode_csv(question_headers, plan["questions"], question_original),
            encode_csv(operation_headers, plan["operation_rows"], operation_original),
        )
        for path, payload in zip(paths, payloads):
            atomic_write_bytes(path, payload)
            if path.read_bytes() != payload:
                raise RuntimeError(f"Write verification failed for {path.name}")
    print(f"Applied. Submissions: {len(plan['submissions'])}; cells: {len(plan['cells'])}")
    print(f"Backup: {backup_dir}")
    return plan


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", required=True, type=Path)
    parser.add_argument("--xlsx", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    data_dir = args.data_dir.expanduser().resolve()
    xlsx = args.xlsx or data_dir / "QualityMatrix" / "Source" / "Test Quality Product Matrix_Curent.xlsx"
    data_dir, xlsx = ensure_local_paths(data_dir, xlsx)
    plan = create_plan(data_dir, xlsx)
    print_preview(plan)
    if args.apply:
        apply_plan(plan)
    else:
        print("Preview only; no files changed.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, ValueError, RuntimeError, csv.Error) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)