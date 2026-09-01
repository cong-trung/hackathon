"""Read/write product-basis and solution CSV/XLSX files, mapping headers to API fields."""

import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import polars as pl

from app.config import DATA_DIR


def value_to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_excel_file(path: Path) -> bool:
    return path.exists() and zipfile.is_zipfile(path)


def get_path(config: dict[str, Any]) -> Path:
    return DATA_DIR / config["file"]


def normalize_header(header: str, config: dict[str, Any]) -> str:
    return config["original_to_fe"].get(value_to_str(header), value_to_str(header))


def original_header(key: str, config: dict[str, Any]) -> str:
    return config["fe_to_original"].get(key, key)


def normalize_row(row: dict[str, Any], config: dict[str, Any]) -> dict[str, str]:
    out = {}
    for key, value in row.items():
        out[normalize_header(key, config)] = value_to_str(value)
    for key in config["fields"]:
        out.setdefault(key, "")
    return out


def denormalize_row(
    row: dict[str, Any], headers: list[str], config: dict[str, Any]
) -> dict[str, str]:
    out = {}
    for key in config["fields"]:
        out[original_header(key, config)] = value_to_str(row.get(key, ""))
    for key, value in row.items():
        if key not in config["fields"]:
            out[original_header(key, config)] = value_to_str(value)
    for header in headers:
        out.setdefault(header, "")
    return out


def read_excel_like(
    path: Path, config: dict[str, Any]
) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    tmp_path = None
    load_path = path
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp_path = Path(tmp.name)
        shutil.copy(path, tmp_path)
        load_path = tmp_path

    try:
        wb = load_workbook(load_path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [
            value_to_str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)
        ]
        rows = []
        for r in range(2, ws.max_row + 1):
            raw = {
                headers[c - 1]: ws.cell(r, c).value for c in range(1, len(headers) + 1)
            }
            if any(value_to_str(v) for v in raw.values()):
                rows.append(normalize_row(raw, config))
        return headers, rows
    finally:
        if tmp_path:
            tmp_path.unlink(missing_ok=True)


def write_excel_like(
    path: Path, rows: list[dict[str, Any]], headers: list[str], config: dict[str, Any]
) -> None:
    from openpyxl import Workbook, load_workbook

    final_headers = [config["fe_to_original"][key] for key in config["fields"]]
    for header in headers:
        if header not in final_headers:
            final_headers.append(header)
    for row in rows:
        for key in row:
            header = original_header(key, config)
            if header not in final_headers:
                final_headers.append(header)

    if path.exists() and is_excel_file(path):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_in:
            tmp_in_path = Path(tmp_in.name)
        shutil.copy(path, tmp_in_path)
        wb = load_workbook(tmp_in_path)
        tmp_in_path.unlink(missing_ok=True)
        ws = wb[wb.sheetnames[0]]
        ws.delete_rows(1, ws.max_row)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

    ws.append(final_headers)
    for row in rows:
        denorm = denormalize_row(row, final_headers, config)
        ws.append([denorm.get(header, "") for header in final_headers])

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_out:
        tmp_out_path = Path(tmp_out.name)
    wb.save(tmp_out_path)
    shutil.copy(tmp_out_path, path)
    tmp_out_path.unlink(missing_ok=True)


def read_csv_file(
    path: Path, config: dict[str, Any]
) -> tuple[list[str], list[dict[str, str]]]:
    if not path.exists():
        return [config["fe_to_original"][key] for key in config["fields"]], []
    lf = pl.scan_csv(path, encoding="utf8-lossy", infer_schema_length=0)
    original_headers = lf.collect_schema().names()
    rename_map = {
        k: v for k, v in config["original_to_fe"].items() if k in original_headers
    }
    df = lf.rename(rename_map).collect()
    rows = []
    for row in df.to_dicts():
        normalized = {k: value_to_str(v) for k, v in row.items()}
        for key in config["fields"]:
            normalized.setdefault(key, "")
        rows.append(normalized)
    return list(original_headers), rows


def write_csv_file(
    path: Path, rows: list[dict[str, Any]], headers: list[str], config: dict[str, Any]
) -> None:
    final_headers = [config["fe_to_original"][key] for key in config["fields"]]
    for header in headers:
        if header not in final_headers:
            final_headers.append(header)
    for row in rows:
        for key in row:
            header = original_header(key, config)
            if header not in final_headers:
                final_headers.append(header)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = [denormalize_row(row, final_headers, config) for row in rows]
    df = pl.DataFrame(
        {h: [r.get(h, "") for r in data] for h in final_headers},
        schema={h: pl.Utf8 for h in final_headers},
    )
    path.write_bytes(b"\xef\xbb\xbf" + df.write_csv().encode("utf-8"))


def read_data(config: dict[str, Any]) -> tuple[list[str], list[dict[str, str]]]:
    path = get_path(config)
    path.parent.mkdir(parents=True, exist_ok=True)
    if is_excel_file(path):
        return read_excel_like(path, config)
    return read_csv_file(path, config)


def write_data(
    config: dict[str, Any], rows: list[dict[str, Any]], headers: list[str]
) -> None:
    path = get_path(config)
    path.parent.mkdir(parents=True, exist_ok=True)
    if is_excel_file(path):
        write_excel_like(path, rows, headers, config)
    else:
        write_csv_file(path, rows, headers, config)


def find_by_key(rows: list[dict[str, str]], key: str, value: str) -> int:
    target = value_to_str(value).lower()
    for idx, row in enumerate(rows):
        if value_to_str(row.get(key, "")).lower() == target:
            return idx
    return -1
