import os
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import polars as pl
from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict, Field

DATA_DIR = Path(os.getenv("QUALITY_MATRIX_DATA_DIR", "data"))
ALLOWED_ORIGINS = [
    x.strip()
    for x in os.getenv(
        "ALLOWED_ORIGINS", "http://localhost:5173,http://localhost:3000"
    ).split(",")
]

app = FastAPI(
    title="Quality Matrix Multi-Module API",
    version="1.0.0",
    description="No-auth GitHub-ready API for STHI, LCBI product basic info and module solutions.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MODULE_CONFIG = {
    "STHI": {
        "file": "ProductBasisInfo_STHI.csv",
        "fields": [
            "product",
            "prodgroup3",
            "module",
            "nsb",
            "mcp",
            "pkg",
            "pkg_1",
            "prd_seg",
            "xvi_tool_type",
            "fab_tech",
            "thermal_tech",
            "sw_platform",
            "lts_ball",
            "fusion_or_apse",
            "socket_type",
        ],
        "original_to_fe": {
            "Product (Die code name)": "product",
            "ProdGroup3": "prodgroup3",
            "Module": "module",
            "NSB? \n(Non-Synergy-Build)": "nsb",
            "NSB?\n(Non-Synergy-Build)": "nsb",
            "NSB? (Non-Synergy-Build)": "nsb",
            "MCP (Multi chip package - products have multiple dies?)": "mcp",
            "Package": "pkg",
            "Package (1)": "pkg_1",
            "Product segment": "prd_seg",
            "xVI tool type": "xvi_tool_type",
            "Fab Technology": "fab_tech",
            "Thermal Technology": "thermal_tech",
            "SW platform": "sw_platform",
            "LTS ball": "lts_ball",
            "Fusion or Apse": "fusion_or_apse",
            "Socket Type": "socket_type",
        },
        "fe_to_original": {
            "product": "Product (Die code name)",
            "prodgroup3": "ProdGroup3",
            "module": "Module",
            "nsb": "NSB? \n(Non-Synergy-Build)",
            "mcp": "MCP (Multi chip package - products have multiple dies?)",
            "pkg": "Package",
            "pkg_1": "Package (1)",
            "prd_seg": "Product segment",
            "xvi_tool_type": "xVI tool type",
            "fab_tech": "Fab Technology",
            "thermal_tech": "Thermal Technology",
            "sw_platform": "SW platform",
            "lts_ball": "LTS ball",
            "fusion_or_apse": "Fusion or Apse",
            "socket_type": "Socket Type",
        },
    },
    "LCBI": {
        "file": "ProductBasicInfo_LCBI.csv",
        "fields": [
            "product",
            "prodgroup3",
            "module",
            "pkg",
            "pkg_1",
            "prd_seg",
            "fab_tech",
        ],
        "original_to_fe": {
            "Product": "product",
            "ProdGroup3": "prodgroup3",
            "Module": "module",
            "Package": "pkg",
            "Package (1)": "pkg_1",
            "Product segment": "prd_seg",
            "Fab Technology": "fab_tech",
        },
        "fe_to_original": {
            "product": "Product",
            "prodgroup3": "ProdGroup3",
            "module": "Module",
            "pkg": "Package",
            "pkg_1": "Package (1)",
            "prd_seg": "Product segment",
            "fab_tech": "Fab Technology",
        },
    },
}

SOLUTION_CONFIG = {
    "file": "Solution.csv",
    "fields": ["id", "module", "solution", "key_summary", "excursion_related"],
    "original_to_fe": {
        "ID": "id",
        "Module": "module",
        "Solution": "solution",
        "Key Summary": "key_summary",
        "Excursion Releated": "excursion_related",
        "Excursion Related": "excursion_related",
    },
    "fe_to_original": {
        "id": "ID",
        "module": "Module",
        "solution": "Solution",
        "key_summary": "Key Summary",
        "excursion_related": "Excursion Releated",
    },
}

FIELD_LABELS = {
    "product": "Product",
    "prodgroup3": "ProdGroup3",
    "module": "Module",
    "nsb": "NSB",
    "mcp": "MCP",
    "pkg": "Package",
    "pkg_1": "Package (1)",
    "prd_seg": "Product Segment",
    "xvi_tool_type": "xVI Tool Type",
    "fab_tech": "Fab Technology",
    "thermal_tech": "Thermal Technology",
    "sw_platform": "SW Platform",
    "lts_ball": "LTS Ball",
    "fusion_or_apse": "Fusion or Apse",
    "socket_type": "Socket Type",
    "id": "ID",
    "solution": "Solution",
    "key_summary": "Key Summary",
    "excursion_related": "Excursion Related",
}

MODULE_SOLUTION_CONFIG = {
    "STHI": {
        "file": "ProductSolutionMatrix_STHI.csv",
        "fields": ["pdid", "solutionid", "status"],
        "original_to_fe": {
            "PdID": "pdid",
            "SolutionID": "solutionid",
            "Status": "status",
        },
        "fe_to_original": {
            "pdid": "PdID",
            "solutionid": "SolutionID",
            "status": "Status",
        },
    },
    "LCBI": {
        "file": "ProductSolutionMatrix_LCBI.csv",
        "fields": ["pdid", "solutionid", "status"],
        "original_to_fe": {
            "PdID": "pdid",
            "SolutionID": "solutionid",
            "Status": "status",
        },
        "fe_to_original": {
            "pdid": "PdID",
            "solutionid": "SolutionID",
            "status": "Status",
        },
    },
}


def value_to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_excel_file(path: Path) -> bool:
    return path.exists() and zipfile.is_zipfile(path)


def get_module_config(module: str) -> dict[str, Any]:
    key = module.upper().strip()
    if key not in MODULE_CONFIG:
        raise HTTPException(status_code=404, detail=f"Unsupported module '{module}'")
    return MODULE_CONFIG[key]


def get_module_solution_config(module: str) -> dict[str, Any]:
    key = module.upper().strip()
    if key not in MODULE_SOLUTION_CONFIG:
        raise HTTPException(status_code=404, detail=f"Unsupported module '{module}'")
    return MODULE_SOLUTION_CONFIG[key]


def get_path(config: dict[str, Any]) -> Path:
    return DATA_DIR / config["file"]


def normalize_header(header: str, config: dict[str, Any]) -> str:
    return config["original_to_fe"].get(value_to_str(header), value_to_str(header))


def original_header(key: str, config: dict[str, Any]) -> str:
    return config["fe_to_original"].get(key, key)


def normalize_row(row: dict[str, Any], config: dict[str, Any]) -> dict[str, str]:
    out = {}
    for key, value in row.items():
        out[normalize_header(key, config)] = value_to_str(value)
    for key in config["fields"]:
        out.setdefault(key, "")
    return out


def denormalize_row(
    row: dict[str, Any], headers: list[str], config: dict[str, Any]
) -> dict[str, str]:
    out = {}
    for key in config["fields"]:
        out[original_header(key, config)] = value_to_str(row.get(key, ""))
    for key, value in row.items():
        if key not in config["fields"]:
            out[original_header(key, config)] = value_to_str(value)
    for header in headers:
        out.setdefault(header, "")
    return out


def read_excel_like(
    path: Path, config: dict[str, Any]
) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    tmp_path = None
    load_path = path
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp_path = Path(tmp.name)
        shutil.copy(path, tmp_path)
        load_path = tmp_path

    try:
        wb = load_workbook(load_path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [
            value_to_str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)
        ]
        rows = []
        for r in range(2, ws.max_row + 1):
            raw = {
                headers[c - 1]: ws.cell(r, c).value for c in range(1, len(headers) + 1)
            }
            if any(value_to_str(v) for v in raw.values()):
                rows.append(normalize_row(raw, config))
        return headers, rows
    finally:
        if tmp_path:
            tmp_path.unlink(missing_ok=True)


def write_excel_like(
    path: Path, rows: list[dict[str, Any]], headers: list[str], config: dict[str, Any]
) -> None:
    from openpyxl import Workbook, load_workbook

    final_headers = [config["fe_to_original"][key] for key in config["fields"]]
    for header in headers:
        if header not in final_headers:
            final_headers.append(header)
    for row in rows:
        for key in row:
            header = original_header(key, config)
            if header not in final_headers:
                final_headers.append(header)

    if path.exists() and is_excel_file(path):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_in:
            tmp_in_path = Path(tmp_in.name)
        shutil.copy(path, tmp_in_path)
        wb = load_workbook(tmp_in_path)
        tmp_in_path.unlink(missing_ok=True)
        ws = wb[wb.sheetnames[0]]
        ws.delete_rows(1, ws.max_row)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

    ws.append(final_headers)
    for row in rows:
        denorm = denormalize_row(row, final_headers, config)
        ws.append([denorm.get(header, "") for header in final_headers])

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_out:
        tmp_out_path = Path(tmp_out.name)
    wb.save(tmp_out_path)
    shutil.copy(tmp_out_path, path)
    tmp_out_path.unlink(missing_ok=True)


def read_csv_file(
    path: Path, config: dict[str, Any]
) -> tuple[list[str], list[dict[str, str]]]:
    if not path.exists():
        return [config["fe_to_original"][key] for key in config["fields"]], []
    lf = pl.scan_csv(path, encoding="utf8-lossy", infer_schema_length=0)
    original_headers = lf.collect_schema().names()
    rename_map = {
        k: v for k, v in config["original_to_fe"].items() if k in original_headers
    }
    df = lf.rename(rename_map).collect()
    rows = []
    for row in df.to_dicts():
        normalized = {k: value_to_str(v) for k, v in row.items()}
        for key in config["fields"]:
            normalized.setdefault(key, "")
        rows.append(normalized)
    return list(original_headers), rows


def write_csv_file(
    path: Path, rows: list[dict[str, Any]], headers: list[str], config: dict[str, Any]
) -> None:
    final_headers = [config["fe_to_original"][key] for key in config["fields"]]
    for header in headers:
        if header not in final_headers:
            final_headers.append(header)
    for row in rows:
        for key in row:
            header = original_header(key, config)
            if header not in final_headers:
                final_headers.append(header)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = [denormalize_row(row, final_headers, config) for row in rows]
    df = pl.DataFrame(
        {h: [r.get(h, "") for r in data] for h in final_headers},
        schema={h: pl.Utf8 for h in final_headers},
    )
    path.write_bytes(b"\xef\xbb\xbf" + df.write_csv().encode("utf-8"))


def read_data(config: dict[str, Any]) -> tuple[list[str], list[dict[str, str]]]:
    path = get_path(config)
    path.parent.mkdir(parents=True, exist_ok=True)
    if is_excel_file(path):
        return read_excel_like(path, config)
    return read_csv_file(path, config)


def write_data(
    config: dict[str, Any], rows: list[dict[str, Any]], headers: list[str]
) -> None:
    path = get_path(config)
    path.parent.mkdir(parents=True, exist_ok=True)
    if is_excel_file(path):
        write_excel_like(path, rows, headers, config)
    else:
        write_csv_file(path, rows, headers, config)


def find_by_key(rows: list[dict[str, str]], key: str, value: str) -> int:
    target = value_to_str(value).lower()
    for idx, row in enumerate(rows):
        if value_to_str(row.get(key, "")).lower() == target:
            return idx
    return -1


def field_schema(config: dict[str, Any]) -> dict[str, Any]:
    fields = []
    for key in config["fields"]:
        label = FIELD_LABELS.get(key, key)
        fields.append(
            {
                "key": key,
                "label": label,
                "type": "string",
                "required": key not in {"key_summary", "excursion_related"},
                "minLength": 1
                if key not in {"key_summary", "excursion_related"}
                else 0,
                "placeholder": label,
                "errorMessage": f"{label} is required"
                if key not in {"key_summary", "excursion_related"}
                else "",
            }
        )
    return {"fields": fields}


class FlexiblePayload(BaseModel):
    model_config = ConfigDict(extra="allow")


class ProductBasisPayload(BaseModel):
    model_config = ConfigDict(extra="allow")
    product: str = Field(..., min_length=1)
    prodgroup3: str = Field(..., min_length=1)
    module: str = Field(..., min_length=1)


class SolutionPayload(BaseModel):
    model_config = ConfigDict(extra="allow")
    module: str = Field(..., min_length=1)
    solution: str = Field(..., min_length=1)
    key_summary: str | None = ""
    excursion_related: str | None = ""


class PdSolutionItem(BaseModel):
    pdid: str = Field(..., min_length=1)
    solutionid: str = Field(..., min_length=1)
    status: str = ""


@app.get("/health")
def health():
    result = {"status": "ok", "modules": {}}
    for module, config in MODULE_CONFIG.items():
        headers, rows = read_data(config)
        result["modules"][module] = {
            "file": str(get_path(config)),
            "file_is_excel_format": is_excel_file(get_path(config)),
            "row_count": len(rows),
            "column_count": len(headers),
        }
    headers, rows = read_data(SOLUTION_CONFIG)
    result["solutions"] = {
        "file": str(get_path(SOLUTION_CONFIG)),
        "row_count": len(rows),
        "column_count": len(headers),
    }
    return result


@app.get("/modules")
def modules():
    return {"items": list(MODULE_CONFIG.keys())}


@app.get("/schema/product-basis/{module}")
def schema_product_basis(module: str):
    return field_schema(get_module_config(module))


@app.get("/schema/solution")
def schema_solution():
    return field_schema(SOLUTION_CONFIG)


@app.get("/product-basis/{module}")
def list_product_basis(
    module: str,
    q: str | None = None,
    limit: int = Query(200, ge=1, le=5000),
    offset: int = Query(0, ge=0),
):
    config = get_module_config(module)
    _headers, rows = read_data(config)
    if q:
        qn = q.lower()
        rows = [row for row in rows if qn in " ".join(row.values()).lower()]
    rows = sorted(rows, key=lambda r: r.get("prodgroup3", "").lower())
    return {
        "module": module.upper(),
        "count": len(rows[offset : offset + limit]),
        "total": len(rows),
        "items": rows[offset : offset + limit],
    }


@app.get("/product-basis/{module}/{prodgroup3}")
def get_product_basis(module: str, prodgroup3: str):
    config = get_module_config(module)
    _headers, rows = read_data(config)
    idx = find_by_key(rows, "prodgroup3", prodgroup3)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Product not found")
    return rows[idx]


@app.post("/product-basis/{module}")
def create_product_basis(module: str, payload: ProductBasisPayload):
    config = get_module_config(module)
    headers, rows = read_data(config)
    data = payload.model_dump()
    if find_by_key(rows, "prodgroup3", str(data.get("prodgroup3", ""))) >= 0:
        raise HTTPException(status_code=409, detail="ProdGroup3 already exists")
    for key in config["fields"]:
        data.setdefault(key, "")
    data["module"] = module.upper()
    rows.append({k: value_to_str(v) for k, v in data.items()})
    write_data(config, rows, headers)
    return {"message": "created", "item": data}


@app.put("/product-basis/{module}/{prodgroup3}")
def update_product_basis(module: str, prodgroup3: str, payload: ProductBasisPayload):
    config = get_module_config(module)
    headers, rows = read_data(config)
    idx = find_by_key(rows, "prodgroup3", prodgroup3)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Product not found")
    data = payload.model_dump()
    rows[idx] = {**rows[idx], **{k: value_to_str(v) for k, v in data.items()}}
    write_data(config, rows, headers)
    return {"message": "updated", "item": rows[idx]}


@app.patch("/product-basis/{module}/{prodgroup3}")
def patch_product_basis(module: str, prodgroup3: str, payload: FlexiblePayload):
    config = get_module_config(module)
    headers, rows = read_data(config)
    idx = find_by_key(rows, "prodgroup3", prodgroup3)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Product not found")
    data = payload.model_dump(exclude_none=True)
    rows[idx] = {**rows[idx], **{k: value_to_str(v) for k, v in data.items()}}
    write_data(config, rows, headers)
    return {"message": "patched", "item": rows[idx]}


@app.delete("/product-basis/{module}/{prodgroup3}")
def delete_product_basis(module: str, prodgroup3: str):
    config = get_module_config(module)
    headers, rows = read_data(config)
    idx = find_by_key(rows, "prodgroup3", prodgroup3)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Product not found")
    removed = rows.pop(idx)
    write_data(config, rows, headers)
    return {"message": "deleted", "item": removed}


@app.get("/solutions")
def list_solutions(
    module: str | None = None,
    q: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
    offset: int = Query(0, ge=0),
):
    _headers, rows = read_data(SOLUTION_CONFIG)
    if module:
        rows = [row for row in rows if row.get("module", "").upper() == module.upper()]
    if q:
        qn = q.lower()
        rows = [row for row in rows if qn in " ".join(row.values()).lower()]
    return {
        "count": len(rows[offset : offset + limit]),
        "total": len(rows),
        "items": rows[offset : offset + limit],
    }


@app.get("/solutions/{solution_id}")
def get_solution(solution_id: str):
    _headers, rows = read_data(SOLUTION_CONFIG)
    idx = find_by_key(rows, "id", solution_id)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Solution not found")
    return rows[idx]


@app.post("/solutions")
def create_solution(payload: SolutionPayload):
    headers, rows = read_data(SOLUTION_CONFIG)
    data = payload.model_dump()
    existing_ids = [
        int(row["id"]) for row in rows if value_to_str(row.get("id", "")).isdigit()
    ]
    data["id"] = str(max(existing_ids, default=0) + 1)
    rows.append({k: value_to_str(v) for k, v in data.items()})
    write_data(SOLUTION_CONFIG, rows, headers)
    return {"message": "created", "item": data}


@app.patch("/solutions/{solution_id}")
def patch_solution(solution_id: str, payload: FlexiblePayload):
    headers, rows = read_data(SOLUTION_CONFIG)
    idx = find_by_key(rows, "id", solution_id)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Solution not found")
    data = payload.model_dump(exclude_none=True)
    rows[idx] = {**rows[idx], **{k: value_to_str(v) for k, v in data.items()}}
    write_data(SOLUTION_CONFIG, rows, headers)
    return {"message": "patched", "item": rows[idx]}


@app.delete("/solutions/{solution_id}")
def delete_solution(solution_id: str):
    headers, rows = read_data(SOLUTION_CONFIG)
    idx = find_by_key(rows, "id", solution_id)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Solution not found")
    removed = rows.pop(idx)
    write_data(SOLUTION_CONFIG, rows, headers)
    return {"message": "deleted", "item": removed}


@app.get("/pdsolutions/{module}")
def list_pdsolutions(module: str, q: str | None = None):
    config = get_module_solution_config(module)
    _headers, rows = read_data(config)

    _sol_headers, sol_rows = read_data(SOLUTION_CONFIG)
    sol_map = {r["id"]: r.get("solution", "") for r in sol_rows}

    if q:
        qn = q.lower()
        rows = [row for row in rows if qn in " ".join(row.values()).lower()]

    grouped: dict[str, list[dict]] = {}
    for row in rows:
        pdid = row.get("pdid", "")
        sid = row.get("solutionid", "")
        grouped.setdefault(pdid, []).append(
            {
                "solutionid": sid,
                "solution": sol_map.get(sid, ""),
                "status": row.get("status", ""),
            }
        )

    for items in grouped.values():
        items.sort(
            key=lambda x: (
                int(x["solutionid"]) if x["solutionid"].isdigit() else x["solutionid"]
            )
        )

    return {"module": module.upper(), "total": len(grouped), "items": grouped}


@app.post("/pdsolutions/{module}")
def update_pdsolutions(module: str, payload: list[PdSolutionItem]):
    config = get_module_solution_config(module)
    headers, rows = read_data(config)
    for item in payload:
        idx = next(
            (
                i
                for i, r in enumerate(rows)
                if r.get("pdid", "").lower() == item.pdid.lower()
                and r.get("solutionid", "").lower() == item.solutionid.lower()
            ),
            -1,
        )
        if idx >= 0:
            rows[idx]["status"] = item.status
        else:
            rows.append(
                {
                    "pdid": item.pdid,
                    "solutionid": item.solutionid,
                    "status": item.status,
                }
            )
    write_data(config, rows, headers)
    return {"message": "updated", "module": module.upper(), "count": len(payload)}
